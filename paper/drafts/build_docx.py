from __future__ import annotations

import re
import shutil
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from docx.text.paragraph import Paragraph


HERE = Path(__file__).resolve().parent
FONT = "Malgun Gothic"


def _blocks(lines: list[str]) -> list[tuple[str, object]]:
    blocks: list[tuple[str, object]] = []
    text_lines: list[tuple[str, bool]] = []

    def flush() -> None:
        if text_lines:
            blocks.append(("paragraph", list(text_lines)))
            text_lines.clear()

    i = 0
    while i < len(lines):
        raw = lines[i]
        line = raw.strip()
        if not line:
            flush()
            i += 1
            continue
        if line.startswith("|"):
            flush()
            table_lines: list[str] = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_lines.append(lines[i].strip())
                i += 1
            rows = []
            for row in table_lines:
                cells = [cell.strip() for cell in row.strip("|").split("|")]
                if not all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells):
                    rows.append(cells)
            if rows:
                blocks.append(("table", rows))
            continue
        if line.startswith("!["):
            flush()
            match = re.fullmatch(r"!\[([^]]*)\]\(([^)]+)\)", line)
            if match:
                blocks.append(("image", (match.group(1), match.group(2))))
            i += 1
            continue
        if line.startswith("## "):
            flush()
            blocks.append(("heading", (2, line[3:].strip())))
            i += 1
            continue
        if line.startswith("### "):
            flush()
            blocks.append(("heading", (3, line[4:].strip())))
            i += 1
            continue
        hard_break = raw.endswith("  ") or line.endswith("<br>")
        line = line.removesuffix("<br>").rstrip()
        text_lines.append((line.rstrip(), hard_break))
        i += 1
    flush()
    return blocks


def _set_run(run, *, size: float = 9, bold: bool = False, italic: bool = False) -> None:
    run.font.name = FONT
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{key}"), FONT)


def _set_paragraph(p, *, size: float = 9, bold: bool = False, after: float = 2) -> None:
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.0
    for run in p.runs:
        _set_run(run, size=size, bold=bold, italic=run.italic)


def _paragraph_after(cursor, parent) -> tuple[Paragraph, object]:
    element = OxmlElement("w:p")
    cursor.addnext(element)
    return Paragraph(element, parent), element


def _paragraph_before(cursor, parent) -> tuple[Paragraph, object]:
    element = OxmlElement("w:p")
    cursor.addprevious(element)
    return Paragraph(element, parent), element


def _table_borders(table) -> None:
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = OxmlElement(f"w:{edge}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "4")
        border.set(qn("w:space"), "0")
        border.set(qn("w:color"), "000000")
        borders.append(border)
    table._tbl.tblPr.append(borders)


def _table_layout(table) -> None:
    # Keep tables within one column; python-docx otherwise uses the full page width.
    widths = (Cm(2.4), Cm(2.0), Cm(2.0), Cm(2.0))
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table_width = table._tbl.tblPr.find(qn("w:tblW"))
    if table_width is None:
        table_width = OxmlElement("w:tblW")
        table._tbl.tblPr.append(table_width)
    table_width.set(qn("w:w"), str(sum(width.twips for width in widths)))
    table_width.set(qn("w:type"), "dxa")
    for column, width in zip(table.columns, widths):
        column.width = width
    for row in table.rows:
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        for cell, width in zip(row.cells, widths):
            cell.width = width
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    table.rows[0]._tr.get_or_add_trPr().append(header)


def _write_text(p: Paragraph, lines: list[tuple[str, bool]], *, size: float, bold: bool = False) -> None:
    for index, (text, hard_break) in enumerate(lines):
        if index and not lines[index - 1][1]:
            p.add_run(" ")
        for token in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", text):
            if not token:
                continue
            is_bold = token.startswith("**") and token.endswith("**")
            is_italic = token.startswith("*") and token.endswith("*") and not is_bold
            value = token[2:-2] if is_bold else token[1:-1] if is_italic else token
            run = p.add_run(value)
            _set_run(run, size=size, bold=bold or is_bold, italic=is_italic)
        if hard_break and index + 1 < len(lines):
            p.add_run().add_break()
    _set_paragraph(p, size=size, bold=bold)


def _insert_blocks(doc, cursor, parent, blocks, image_dir: Path, *, before: bool = False):
    anchor = cursor
    for kind, value in blocks:
        if kind == "table":
            rows = value
            table = doc.add_table(rows=len(rows), cols=max(map(len, rows)))
            table.autofit = True
            _table_borders(table)
            if len(rows[0]) == 4:
                _table_layout(table)
            element = table._tbl
            element.getparent().remove(element)
            anchor.addprevious(element) if before else cursor.addnext(element)
            if not before:
                cursor = element
            for row_index, cells in enumerate(rows):
                for column_index, text in enumerate(cells):
                    cell = table.cell(row_index, column_index)
                    cell.text = ""
                    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    p = cell.paragraphs[0]
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if column_index else WD_ALIGN_PARAGRAPH.LEFT
                    p.paragraph_format.space_after = Pt(0)
                    p.paragraph_format.line_spacing = 1.0
                    for part_index, part in enumerate(text.replace("<br>", "\n").splitlines()):
                        if part_index:
                            p.add_run().add_break()
                        run = p.add_run(part)
                        _set_run(run, size=9, bold=row_index == 0)
            continue

        p, element = (
            _paragraph_before(anchor, parent) if before else _paragraph_after(cursor, parent)
        )
        if not before:
            cursor = element
        if kind == "image":
            alt, relative_path = value
            image_path = (image_dir / relative_path).resolve()
            if not image_path.is_file():
                raise FileNotFoundError(image_path)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            shape = p.add_run().add_picture(str(image_path), width=Cm(8.4))
            shape._inline.docPr.set("descr", alt)
        elif kind == "heading":
            level, text = value
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(text)
            _set_run(run, size=10 if level == 2 else 9, bold=True)
            _set_paragraph(p, size=10 if level == 2 else 9, bold=True, after=2)
            p.paragraph_format.keep_with_next = True
        else:
            lines = value
            text = " ".join(part for part, _ in lines).strip()
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            if text.startswith("그림 "):
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                _write_text(p, lines, size=9)
            elif text.startswith("표 "):
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                _write_text(p, lines, size=9, bold=True)
                p.paragraph_format.keep_with_next = True
            elif re.match(r"\[\d+\]\s", text):
                _write_text(p, lines, size=7.5)
                p.paragraph_format.space_after = Pt(0)
            else:
                _write_text(p, lines, size=9)
    return cursor


def build(round_number: str) -> Path:
    source = HERE / f"round{round_number}_draft.md"
    output = HERE / f"round{round_number}_draft.docx"
    template = HERE.parent / "references" / "forms" / "kiise_review_form.docx"
    if not source.is_file():
        raise FileNotFoundError(source)
    shutil.copy2(template, output)
    doc = Document(output)

    body = doc._element.body
    section_breaks = [
        child for child in body
        if child.tag == qn("w:p") and child.find(f".//{qn('w:sectPr')}") is not None
    ]
    if not section_breaks:
        raise RuntimeError("The KIISE review template has no title/body section break.")
    title_break = section_breaks[0]
    container = doc._body
    for child in list(body):
        if child.tag == qn("w:sectPr") or child in section_breaks:
            continue
        body.remove(child)
    for child in section_breaks:
        for node in list(child):
            if node.tag != qn("w:pPr"):
                child.remove(node)

    lines = source.read_text(encoding="utf-8").splitlines()
    title = next((line[2:].strip() for line in lines if line.startswith("# ")), None)
    english = next((line.strip().strip("*") for line in lines if line.startswith("*")), None)
    body_start = next((i for i, line in enumerate(lines) if line.startswith("## 1. ")), None)
    if not title or not english or body_start is None:
        raise ValueError("Draft must contain Korean and English titles and a '## 1.' body heading.")
    abstract_start = next((i for i, line in enumerate(lines) if line.startswith("## 요약")), None)
    if abstract_start is None or abstract_start >= body_start:
        raise ValueError("Draft must contain a summary before section 1.")

    p, _ = _paragraph_before(title_break, container)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(title)
    _set_run(run, size=13, bold=True)
    _set_paragraph(p, size=13, bold=True, after=1)

    p, _ = _paragraph_before(title_break, container)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(english)
    _set_run(run, size=9, italic=True)
    _set_paragraph(p, size=9, after=4)

    abstract_blocks = _blocks(lines[abstract_start:body_start])
    _insert_blocks(doc, title_break, container, abstract_blocks, HERE, before=True)

    body_cursor = title_break
    body_blocks = _blocks(lines[body_start:])
    _insert_blocks(doc, body_cursor, container, body_blocks, HERE)

    # Review copy must not inherit an author identity from the sample form's metadata.
    properties = doc.core_properties
    properties.author = ""
    properties.last_modified_by = ""
    properties.subject = ""
    properties.comments = ""
    properties.keywords = ""
    properties.title = title
    doc.save(output)
    return output


if __name__ == "__main__":
    if len(sys.argv) != 2 or not sys.argv[1].isdigit():
        raise SystemExit("Usage: python build_docx.py <round_number>")
    print(build(sys.argv[1]))
