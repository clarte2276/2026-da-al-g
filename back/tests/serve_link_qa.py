"""Run an isolated manual-QA API: uv run python tests/serve_link_qa.py.

Use VITE_API_BASE_URL=http://127.0.0.1:18080 for the frontend.
Generated samples and the database live in a fresh OS temporary directory.
"""
import os
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path


def main():
    if "--ui" in sys.argv:
        os.environ["VITE_API_BASE_URL"] = "http://127.0.0.1:18080"
        frontend = Path(__file__).resolve().parents[2] / "front" / "link-generator"
        subprocess.run(["npm.cmd" if os.name == "nt" else "npm", "run", "dev", "--",
                        "--host", "127.0.0.1", "--port", "15173", "--strictPort"],
                       cwd=frontend, check=True)
        return
    root = Path(tempfile.mkdtemp(prefix="daalgi-link-qa-"))
    documents = root / "documents"
    documents.mkdir()
    os.environ.update({
        "DATABASE_URL": f"sqlite:///{(root / 'qa.db').as_posix()}",
        "STORAGE_ROOT": str(root / "storage"),
        "LOCAL_DOCUMENT_ROOTS": str(documents),
        "OPENAI_API_KEY": "", "EMBEDDING_PROVIDER": "hash",
    })
    sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
    import uvicorn
    from pptx import Presentation
    from pptx.util import Inches
    from pypdf import PdfWriter
    from test_document_links import sample_word

    sample_word(documents / "01-text.docx")
    pdf = PdfWriter()
    for _ in range(3):
        pdf.add_blank_page(595, 842)
    pdf.write(documents / "02-pages.pdf")
    presentation = Presentation()
    for number in range(1, 4):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1)).text = f"Sample slide {number}"
    presentation.save(documents / "03-slides.pptx")
    raw = Path(__file__).resolve().parents[2] / "data_raw"
    hwpx = next(raw.rglob("*.hwpx"), None)
    hwp = next(raw.rglob("*.hwp"), None)
    for path in (hwpx, hwp):
        if path:
            shutil.copyfile(path, documents / f"sample{path.suffix}")
    print(f"QA data: {root}", flush=True)
    uvicorn.run("app.main:app", host="127.0.0.1", port=18080)


if __name__ == "__main__":
    main()
