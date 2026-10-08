from datetime import date
from io import BytesIO
from pathlib import Path

import pytest

from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import Workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import Session
from sqlalchemy.pool import StaticPool

from app import auth, duty, management
from app.config import Settings
from app.db import Base, get_db


def test_admin_account_and_personal_data_lifecycle(tmp_path, monkeypatch):
    engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    Base.metadata.create_all(engine)
    settings = Settings(_env_file=None, storage_root=tmp_path, bootstrap_admin_username="admin",
                        bootstrap_admin_password="safe-password", enable_test_account=False)
    monkeypatch.setattr(auth, "get_settings", lambda: settings)
    with Session(engine, expire_on_commit=False) as db:
        auth.ensure_bootstrap_accounts(db)
        app = FastAPI()
        app.include_router(auth.router)
        app.include_router(management.router)
        app.dependency_overrides[get_db] = lambda: db
        with TestClient(app) as client:
            login = lambda username, password: client.post("/api/auth/login", json={"username": username, "password": password}).json()["access_token"]
            admin = {"Authorization": f"Bearer {login('admin', 'safe-password')}"}
            assert client.get("/api/admin/users", headers=admin).status_code == 200
            created = client.post("/api/admin/users", headers=admin, json={"username": "21714093", "display_name": "Driver"})
            assert created.status_code == 201
            user_id = created.json()["user"]["id"]
            user = {"Authorization": f"Bearer {login('21714093', created.json()['initial_password'])}"}
            assert client.get("/api/me/conversations", headers=user).status_code == 403
            assert client.post("/api/auth/password", headers=user,
                               json={"current_password": created.json()["initial_password"],
                                     "new_password": "password-123"}).status_code == 200
            conversation = {"id": "old-1", "title": "Question", "createdAt": "2026-06-01T00:00:00Z", "messages": [{"isUser": True, "text": "Hello"}]}
            bookmark = {"regulation": "Rule", "chapter": "1", "version": ""}
            payload = {"conversations": [conversation], "bookmarks": [bookmark]}
            assert client.post("/api/me/import", headers=user, json=payload).status_code == 200
            assert client.post("/api/me/import", headers=user, json=payload).status_code == 200
            assert len(client.get("/api/me/conversations", headers=user).json()) == 1
            assert len(client.get("/api/me/bookmarks", headers=user).json()) == 1
            assert client.get(f"/api/admin/users/{user_id}/bookmarks", headers=admin).status_code == 404
            assert client.get(f"/api/admin/users/{user_id}/conversations", headers=admin).status_code == 404
            assert client.delete(f"/api/admin/users/{user_id}/bookmarks/old-1", headers=admin).status_code == 404
            assert client.delete(f"/api/admin/users/{user_id}/conversations/old-1", headers=admin).status_code == 404
            assert client.get("/api/admin/feedback", headers=admin).status_code == 404
            assert client.patch(f"/api/admin/users/{user_id}", headers=admin, json={"is_active": False}).status_code == 200
            assert client.get("/api/me/bookmarks", headers=user).status_code == 401


def test_duty_parser_reads_selected_month_only():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "교번기관사 근무계획"
    for day in range(1, 29):
        sheet.cell(2, day + 27, date(2026, 2, day))
        sheet.cell(7, day + 27, "61" if day == 1 else "비")
    sheet.cell(7, 1, 1)
    sheet.cell(7, 2, "기관사")
    content = BytesIO()
    workbook.save(content)
    parsed = management.parse_duty_xlsx(content.getvalue(), "2026-02", "duty.xlsx")
    assert len(parsed["dates"]) == 28
    assert parsed["drivers"][0]["days"]["2026-02-01"] == {"code": "61", "type": "night", "turn": "61"}
    assert parsed["turns"] == {}


DUTY_DIR = Path(__file__).resolve().parents[2] / "data" / "근무계획"
ROUTE_FILE = DUTY_DIR / "240801 개정 행로표_수색+신내.xlsx"
OCT_FILE = DUTY_DIR / "2026년 10월 기관사 근무계획.xlsx"
real_data = pytest.mark.skipif(not (ROUTE_FILE.exists() and OCT_FILE.exists()), reason="data/근무계획 missing")


def test_duty_code_classification():
    cases = {"비": "off", "70~": "off", "대62~": "off", "휴19": "rest", "운휴23": "rest", "대휴3": "rest",
             "휴75~": "rest", "대5": "standby", "대62": "standby", "지정[주]": "designated",
             "지정(야)": "designated", "보완교육": "designated", "7": "day", "61": "night",
             "김대현-휴직(1/1~)": "unknown"}
    assert {code: duty.classify(code) for code in cases} == cases


@real_data
def test_real_route_and_october_duty_merge():
    routes = duty.parse_route_xlsx(ROUTE_FILE.read_bytes(), ROUTE_FILE.name)
    assert routes["effective"] == "2024-08-01"
    night = routes["turns"]["평휴:61"]
    assert (night["depart"], night["arrive"], night["nextDay"]) == ("18:58", "08:03", True)
    assert night["workMinutes"] > 0 and night["segments"][0]["kind"] == "run"
    day = routes["turns"]["평평:1"]
    assert (day["depart"], day["arrive"], day["nextDay"]) == ("06:28", "15:35", False)

    parsed = duty.parse_duty_xlsx(OCT_FILE.read_bytes(), "2026-10", OCT_FILE.name)
    assert parsed["office"] == "수색승무사업소" and len(parsed["dates"]) == 31
    assert parsed["dayTypes"]["2026-10-03"] == "휴" and parsed["dayTypes"]["2026-11-01"] == "휴"
    merged = duty.merge_routes(parsed, routes)
    assert merged["turnDetailIsDummy"] is False
    first = merged["drivers"][0]
    assert first["empNo"] == "21714093"
    assert first["days"]["2026-10-01"]["turnKey"] == "평평:70"
    assert first["days"]["2026-10-02"] == {"code": "70~", "type": "off", "turn": None, "dayType": "평",
                                           "turnKey": None, "continues": "평평:70"}
    assert first["days"]["2026-10-08"]["turnKey"] == "평평:4"
    keys = {d[k] for dr in merged["drivers"] for d in dr["days"].values() for k in ("turnKey", "continues")
            if d.get(k)}
    assert keys and keys <= merged["turns"].keys()
    assert duty.merge_routes(parsed, None)["turnDetailIsDummy"] is True


@real_data
def test_duty_route_upload_endpoint(tmp_path, monkeypatch):
    engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    Base.metadata.create_all(engine)
    settings = Settings(_env_file=None, storage_root=tmp_path, bootstrap_admin_username="admin",
                        bootstrap_admin_password="safe-password", enable_test_account=False)
    monkeypatch.setattr(auth, "get_settings", lambda: settings)
    with Session(engine, expire_on_commit=False) as db:
        auth.ensure_bootstrap_accounts(db)
        app = FastAPI()
        app.include_router(auth.router)
        app.include_router(management.router)
        app.dependency_overrides[get_db] = lambda: db
        with TestClient(app) as client:
            token = client.post("/api/auth/login", json={"username": "admin", "password": "safe-password"}).json()["access_token"]
            admin = {"Authorization": f"Bearer {token}"}
            upload = lambda url, path: client.post(url, headers=admin, files={"file": (path.name, path.read_bytes())})
            assert upload("/api/admin/duty/2026-10", OCT_FILE).status_code == 200
            assert client.get("/api/duty/2026-10", headers=admin).json()["turnDetailIsDummy"] is True
            assert upload("/api/admin/duty-routes", ROUTE_FILE).status_code == 200
            assert client.get("/api/admin/duty-routes", headers=admin).json()[0]["effective_date"] == "2024-08-01"
            data = client.get("/api/duty/2026-10", headers=admin).json()
            assert data["turnDetailIsDummy"] is False and "평평:70" in data["turns"]
