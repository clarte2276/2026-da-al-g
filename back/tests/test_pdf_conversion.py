from pathlib import Path

import pytest
from pptx import Presentation
from pptx.util import Inches
from pypdf import PdfReader

from app.config import Settings
from app.services.pdf_conversion import PdfConversionError, PdfConversionService


def test_pdf_conversion_service_creates_one_pdf_page_per_slide(tmp_path: Path) -> None:
    source = tmp_path / "sample.pptx"
    presentation = Presentation()
    for index in range(2):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
        box.text = f"테스트 슬라이드 {index + 1}"
    presentation.save(source)

    output, converter = PdfConversionService(
        Settings(storage_root=tmp_path / "storage", allow_approximate_pdf_fallback=True)
    ).convert(source, "test-cache-key")

    assert converter in {"libreoffice", "reportlab-fallback"}
    assert output.is_file()
    assert len(PdfReader(output).pages) == 2


def test_pdf_conversion_does_not_silently_use_approximate_renderer(tmp_path: Path) -> None:
    source = tmp_path / "sample.pptx"
    presentation = Presentation()
    presentation.slides.add_slide(presentation.slide_layouts[6])
    presentation.save(source)

    service = PdfConversionService(Settings(storage_root=tmp_path / "storage"))
    service._find_converter = lambda: None  # type: ignore[method-assign]

    with pytest.raises(PdfConversionError, match="LibreOffice를 찾을 수 없습니다"):
        service.convert(source, "missing-converter")


def test_render_page_and_find_page_use_the_pdf_rendition(tmp_path: Path) -> None:
    from reportlab.pdfgen.canvas import Canvas

    source = tmp_path / "sample.pdf"
    canvas = Canvas(str(source))
    canvas.drawString(72, 720, "first page")
    canvas.showPage()
    canvas.drawString(72, 720, "emergency brake procedure")
    canvas.showPage()
    canvas.save()

    service = PdfConversionService(Settings(storage_root=tmp_path / "storage"))
    assert service.find_page(source, "key", "emergency brake") == 2
    assert service.find_page(source, "key", "없는 문장입니다") is None

    image = service.render_page(source, "key", 2, 400)
    assert image.is_file() and image.stat().st_size > 0
    assert service.render_page(source, "key", 2, 400) == image
    with pytest.raises(PdfConversionError):
        service.render_page(source, "key", 99, 400)
