import hashlib

import pytest
from docx import Document as WordDocument
from fastapi import FastAPI
from fastapi.testclient import TestClient
from pptx import Presentation
from pypdf import PdfWriter
from sqlalchemy import create_engine, delete, inspect, select
from sqlalchemy.orm import Session
from sqlalchemy.pool import StaticPool

from app.config import Settings
from app.db import Base, get_db
from app.links import content_indexer, router
from app.models import (
    Document,
    DocumentContent,
    DocumentLink,
    DocumentVersion,
    Fragment,
    KnowledgeEdge,
)
from app.parsers import ParserError, get_parser
from app.services import ingestion as ingestion_module
from app.services.document_content import utf16_length, utf16_slice
from app.services.embedding import HashEmbeddingProvider
from app.services.ingestion import IngestionService
from app.services.rag import GraphRAGService


@pytest.fixture
def workspace(tmp_path):
    engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    Base.metadata.create_all(engine)
    settings = Settings(_env_file=None, storage_root=tmp_path / "store",
                        embedding_dimensions=64, openai_api_key=None)
    provider = HashEmbeddingProvider(64)
    service = IngestionService(settings, provider)
    with Session(engine, expire_on_commit=False) as db:
        app = FastAPI()
        app.include_router(router)
        app.dependency_overrides[get_db] = lambda: db
        app.dependency_overrides[content_indexer] = lambda: service
        with TestClient(app) as client:
            yield db, client, service, GraphRAGService(settings, provider), tmp_path
    engine.dispose()


def ingest(service, db, path):
    document, version = service.ingest_new_file(db, path, path.name)
    db.commit()
    return document, version


def text_selection(version, text, exact, occurrence=0):
    start = -1
    for _ in range(occurrence + 1):
        start = text.index(exact, start + 1)
    offset = utf16_length(text[:start])
    return {"kind": "text", "version_id": version.id, "start": offset,
            "end": offset + utf16_length(exact), "exact": exact}


def sample_word(path):
    word = WordDocument()
    word.add_paragraph("ATC 고장 조치 😀 반복 구절")
    word.add_paragraph("중간 문장")
    table = word.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "왼쪽 셀"
    table.cell(0, 1).text = "오른쪽 셀"
    word.add_paragraph("반복 구절 뒤쪽 " + "긴 문서 내용 " * 800 + " 마지막 문장")
    word.save(path)


def test_text_selection_is_complete_exact_and_independent_of_chunks(workspace, monkeypatch):
    db, client, service, _rag, tmp = workspace
    path = tmp / "sample.docx"
    sample_word(path)
    document, version = ingest(service, db, path)
    response = client.get(f"/api/documents/{document.id}/content", params={"version_id": version.id})
    assert response.status_code == 200
    text = response.json()["text"]
    assert text.endswith("마지막 문장")
    assert text.count("왼쪽 셀") == 1
    assert text.index("중간 문장") < text.index("왼쪽 셀\t오른쪽 셀") < text.index("뒤쪽")
    source = text_selection(version, text, "조치 😀 반복 구절\n\n중간")
    target = text_selection(version, text, "반복 구절", occurrence=1)
    assert utf16_slice(text, source["start"], source["end"]) == source["exact"]
    payload = {"source_selection": source, "target_selection": target}
    response = client.post("/api/links", json=payload)
    assert response.status_code == 201, response.text
    link = response.json()
    assert len(client.get("/api/links").json()) == 1
    assert client.post(f"/api/links/{link['id']}/approve", json={"actor": "tester"}).status_code == 200
    fragments = list(db.scalars(select(Fragment).where(Fragment.version_id == version.id)))
    assert len(fragments) > 2
    # Regenerate indexing rows: the source text and human link must not change.
    old_ids = {f.id for f in fragments}
    db.execute(delete(Fragment).where(Fragment.version_id == version.id))
    original_chunks = ingestion_module._chunks
    monkeypatch.setattr(ingestion_module, "_chunks", lambda fragment: original_chunks(
        fragment, max_chars=640, overlap=80,
    ))
    service._persist_parsed(db, version, get_parser(path).parse(path))
    db.commit()
    reindexed = list(db.scalars(select(Fragment)))
    assert len(reindexed) > len(fragments)
    assert old_ids.isdisjoint({f.id for f in reindexed})
    for fragment in reindexed:
        assert utf16_slice(text, fragment.metadata_json["text_start"],
                           fragment.metadata_json["text_end"]) == fragment.text
    saved = client.get(f"/api/links/{link['id']}").json()
    assert saved["status"] == "approved"
    assert saved["source_selection"]["ranges"][0]["exact"] == source["exact"]
    assert saved["target_selection"]["ranges"][0]["start"] == target["start"]
    response = client.put(f"/api/links/{link['id']}", json={**payload, "note": "수정"})
    assert response.json()["status"] == "draft"
    assert response.json()["approved_at"] is None
    # Original source mutation cannot alter the version's preserved bytes.
    path.write_bytes(b"changed")
    assert client.get(f"/api/documents/{document.id}/content").json()["text"] == text


def test_page_groups_include_empty_pages_and_expand_only_after_approval(workspace):
    db, client, service, rag, tmp = workspace
    path = tmp / "sample.docx"
    sample_word(path)
    _, version = ingest(service, db, path)
    text = db.get(DocumentContent, version.id).text
    source = text_selection(version, text, "ATC 고장 조치")
    pdf_path = tmp / "blank.pdf"
    pdf = PdfWriter()
    for _ in range(3):
        pdf.add_blank_page(595, 842)
    pdf.write(pdf_path)
    document, pdf_version = ingest(service, db, pdf_path)
    content = client.get(f"/api/documents/{document.id}/content").json()
    assert [p["number"] for p in content["pages"]] == [1, 2, 3]
    target = {"kind": "pages", "version_id": pdf_version.id, "pages": [3, 1, 3]}
    for invalid_pages in ([], [0], [4]):
        response = client.post("/api/links", json={"source_selection": source,
            "target_selection": {**target, "pages": invalid_pages}})
        assert response.status_code == 422
    link = client.post("/api/links", json={"source_selection": source, "target_selection": target}).json()
    assert link["target_selection"]["pages"] == [1, 3]
    assert not any(e.link_id == link["id"] for e in rag.retrieve(db, "ATC 고장 조치", top_k=1))
    client.post(f"/api/links/{link['id']}/approve", json={"actor": "tester"})
    expanded = rag.retrieve(db, "ATC 고장 조치", top_k=1)
    evidence = next(e for e in expanded if e.link_id == link["id"])
    assert evidence.selection["pages"] == [1, 3]
    assert evidence.document_id == document.id
    assert "텍스트 없는 페이지" in evidence.fragment.text
    pptx_path = tmp / "blank.pptx"
    pptx = Presentation()
    for _ in range(2):
        pptx.slides.add_slide(pptx.slide_layouts[6])
    pptx.save(pptx_path)
    _, slide_version = ingest(service, db, pptx_path)
    pages_link = client.post("/api/links", json={
        "source_selection": link["target_selection"],
        "target_selection": {"kind": "pages", "version_id": slide_version.id, "pages": [2]},
    }).json()
    client.post(f"/api/links/{pages_link['id']}/approve", json={"actor": "tester"})
    assert any(e.link_id == pages_link["id"] and e.hop == 2
               for e in rag.retrieve(db, "ATC 고장 조치", top_k=2, max_hops=2))
    client.post(f"/api/links/{link['id']}/reject", json={"actor": "tester"})
    assert not any(e.link_id == link["id"] for e in rag.retrieve(db, "ATC 고장 조치", top_k=1))


def test_invalid_selections_and_legacy_backfill_preserve_data(workspace):
    db, client, _service, _rag, tmp = workspace
    path = tmp / "legacy.docx"
    sample_word(path)
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    document = Document(filename=path.name, sha256=digest, source_path=str(path), status="ready")
    db.add(document)
    db.flush()
    version = DocumentVersion(document_id=document.id, version_number=1, sha256=digest,
                              storage_path=str(path), status="completed")
    db.add(version)
    db.flush()
    for key in ("a", "b"):
        db.add(Fragment(version_id=version.id, stable_key=key, kind="paragraph", ordinal=0,
                        text=key, content_hash=key))
    db.flush()
    fragments = list(db.scalars(select(Fragment)))
    old = KnowledgeEdge(source_fragment_id=fragments[0].id, target_fragment_id=fragments[1].id,
                        relation_type="RELATED", status="approved")
    db.add(old)
    db.commit()
    text = client.get(f"/api/documents/{document.id}/content").json()["text"]
    assert db.get(KnowledgeEdge, old.id).status == "approved"
    assert db.scalar(select(Fragment.id).where(Fragment.stable_key.like("canonical%")))
    assert len(list(db.scalars(select(DocumentLink)))) == 0
    source = text_selection(version, text, "ATC 고장 조치")
    target = text_selection(version, text, "마지막 문장")
    for invalid in ({**source, "exact": "다른 문장"}, {**source, "end": 999999},
                    {**source, "version_id": "missing"},
                    {"kind": "pages", "version_id": version.id, "pages": [1]}):
        assert client.post("/api/links", json={"source_selection": invalid,
                                              "target_selection": target}).status_code in {404, 422}
    emoji = utf16_length(text[:text.index("😀")])
    with pytest.raises(ValueError):
        utf16_slice(text, emoji, emoji + 1)
    path.write_bytes(b"changed")
    assert client.get(f"/api/documents/{document.id}/content").status_code == 409
    path.unlink()
    assert client.get(f"/api/documents/{document.id}/content").status_code == 404

def test_many_passages_link_to_many_pages_in_one_link(workspace):
    db, client, service, rag, tmp = workspace
    path = tmp / "sample.docx"
    sample_word(path)
    _, version = ingest(service, db, path)
    text = db.get(DocumentContent, version.id).text
    first = text_selection(version, text, "ATC 고장 조치")
    second = text_selection(version, text, "마지막 문장")
    source = {"kind": "text", "version_id": version.id,
              "ranges": [{k: part[k] for k in ("start", "end", "exact")}
                         for part in (second, first)]}
    pdf_path = tmp / "blank.pdf"
    pdf = PdfWriter()
    for _ in range(3):
        pdf.add_blank_page(595, 842)
    pdf.write(pdf_path)
    document, pdf_version = ingest(service, db, pdf_path)
    target = {"kind": "pages", "version_id": pdf_version.id, "pages": [3, 1]}

    overlapping = {**source, "ranges": [source["ranges"][1],
                                        {**first, "end": first["end"] - 1,
                                         "exact": first["exact"][:-1]}]}
    assert client.post("/api/links", json={"source_selection": overlapping,
                                           "target_selection": target}).status_code == 422

    link = client.post("/api/links", json={"source_selection": source,
                                           "target_selection": target}).json()
    stored = link["source_selection"]["ranges"]
    assert [part["exact"] for part in stored] == [first["exact"], second["exact"]]
    assert link["target_selection"]["pages"] == [1, 3]
    assert len(list(db.scalars(select(DocumentLink)))) == 1

    client.post(f"/api/links/{link['id']}/approve", json={"actor": "tester"})
    # Either passage reaches the same page group, and the group reaches both passages back.
    for question in ("ATC 고장 조치", "마지막 문장"):
        evidence = next(e for e in rag.retrieve(db, question, top_k=2) if e.link_id == link["id"])
        assert evidence.selection["pages"] == [1, 3]
        assert evidence.document_id == document.id

def test_docx_rejects_invalid_file(tmp_path):
    path = tmp_path / "invalid.docx"
    path.write_bytes(b"not a word document")
    with pytest.raises(ParserError):
        get_parser(path).parse(path)


def test_empty_docx_never_sends_empty_embedding_input(workspace, monkeypatch):
    db, client, service, _rag, tmp = workspace
    def unexpected_embedding(_texts):
        raise AssertionError("An empty document must not call the embedding provider")
    monkeypatch.setattr(service.embedding_provider, "embed", unexpected_embedding)
    path = tmp / "empty.docx"
    WordDocument().save(path)
    document, _version = ingest(service, db, path)
    response = client.get(f"/api/documents/{document.id}/content")
    assert response.status_code == 200
    assert response.json()["text"] == ""


def test_additive_tables_preserve_legacy_rows():
    engine = create_engine("sqlite://")
    old_tables = [table for table in Base.metadata.sorted_tables
                  if table.name not in {"document_links", "document_contents"}]
    Base.metadata.create_all(engine, tables=old_tables)
    with Session(engine) as db:
        db.add(Document(id="existing", filename="old.hwp", sha256="a" * 64, status="ready"))
        db.commit()
    assert "document_links" not in inspect(engine).get_table_names()
    Base.metadata.create_all(engine)
    Base.metadata.create_all(engine)
    with Session(engine) as db:
        assert db.get(Document, "existing").filename == "old.hwp"
        assert len(list(db.scalars(select(DocumentLink)))) == 0
    assert {"document_links", "document_contents"} <= set(inspect(engine).get_table_names())
    engine.dispose()
