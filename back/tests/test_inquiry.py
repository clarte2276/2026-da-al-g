from io import BytesIO

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from sqlalchemy import create_engine, select
from sqlalchemy.orm import Session
from sqlalchemy.pool import StaticPool

from app import auth, management
from app.config import Settings
from app.db import Base, get_db
from app.models import AuditEvent, Inquiry, User


@pytest.fixture
def api(tmp_path, monkeypatch):
    engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    Base.metadata.create_all(engine)
    settings = Settings(_env_file=None, storage_root=tmp_path, bootstrap_admin_username="admin",
                        bootstrap_admin_password="safe-password", enable_test_account=False)
    monkeypatch.setattr(auth, "get_settings", lambda: settings)
    with Session(engine, expire_on_commit=False) as db:
        auth.ensure_bootstrap_accounts(db)
        app = FastAPI()
        app.include_router(auth.router)
        app.include_router(management.router)
        app.dependency_overrides[get_db] = lambda: db
        with TestClient(app) as client:
            token = client.post("/api/auth/login", json={"username": "admin", "password": "safe-password"}).json()["access_token"]
            yield client, db, {"Authorization": f"Bearer {token}"}
    engine.dispose()


def _create_user(client, admin, username="21714093"):
    response = client.post("/api/admin/users", headers=admin,
                           json={"username": username, "display_name": "직원"})
    assert response.status_code == 201
    return response.json()


def _login(client, username, password):
    response = client.post("/api/auth/login", json={"username": username, "password": password})
    assert response.status_code == 200
    return {"Authorization": f"Bearer {response.json()['access_token']}"}


def test_first_password_change_and_reset(api):
    client, db, admin = api
    created = _create_user(client, admin)
    user = created["user"]
    initial_password = created["initial_password"]
    assert len(initial_password) >= 10
    assert user["must_change_password"] is True
    listed = client.get("/api/admin/users", headers=admin)
    assert listed.status_code == 200 and listed.json()[0]["must_change_password"] is True
    stored_user = db.get(User, user["id"])
    assert initial_password not in stored_user.password_hash

    member = _login(client, user["username"], initial_password)
    assert client.get("/api/auth/me", headers=member).json()["must_change_password"] is True
    denied = client.get("/api/me/inquiries", headers=member)
    assert denied.status_code == 403 and denied.json()["detail"] == "password_change_required"
    assert client.post("/api/auth/logout", headers=member).status_code == 200

    member = _login(client, user["username"], initial_password)
    changed = client.post("/api/auth/password", headers=member,
                          json={"current_password": initial_password, "new_password": "member-password"})
    assert changed.status_code == 200
    assert client.get("/api/auth/me", headers=member).json()["must_change_password"] is False
    assert client.get("/api/me/inquiries", headers=member).status_code == 200

    reset = client.patch(f"/api/admin/users/{user['id']}", headers=admin, json={"reset_password": True})
    assert reset.status_code == 200
    assert reset.json()["user"]["must_change_password"] is True
    assert len(reset.json()["initial_password"]) >= 10
    assert client.get("/api/auth/me", headers=member).status_code == 401


def test_user_xlsx_template_import_and_credentials(api):
    client, db, admin = api
    template = client.get("/api/admin/users/template.xlsx", headers=admin)
    assert template.status_code == 200
    workbook = load_workbook(BytesIO(template.content), read_only=True, data_only=True)
    assert list(workbook.active.values)[0] == ("사번", "이름")
    workbook.close()

    existing = _create_user(client, admin, "998877")
    sheet = Workbook()
    page = sheet.active
    page.append(["사번", "이름"])
    page.append(["12345678", "가나다"])
    page.append(["12345678", "중복"])
    page.append(["bad-id", "잘못된 사번"])
    page.append(["998877", "기존 사용자"])
    content = BytesIO()
    sheet.save(content)
    response = client.post("/api/admin/users/import", headers=admin,
                           files={"file": ("members.xlsx", content.getvalue())})
    assert response.status_code == 200
    result = response.json()
    assert [row["username"] for row in result["created"]] == ["12345678"]
    assert len(result["created"][0]["initial_password"]) >= 10
    assert len(result["skipped"]) == 3
    assert all(item["reason"] for item in result["skipped"])
    imported = db.scalar(select(User).where(User.username == "12345678"))
    assert imported.must_change_password is True
    assert imported.password_hash != result["created"][0]["initial_password"]
    assert existing["user"]["must_change_password"] is True

    credentials = client.post("/api/admin/users/credentials.xlsx", headers=admin,
                              json={"rows": result["created"]})
    assert credentials.status_code == 200
    workbook = load_workbook(BytesIO(credentials.content), read_only=True, data_only=True)
    rows = list(workbook.active.values)
    assert rows == [("사번", "이름", "초기 비밀번호"),
                    ("12345678", "가나다", result["created"][0]["initial_password"])]
    workbook.close()


def test_user_and_admin_inquiry_replies_and_feedback(api):
    client, db, admin = api
    created = _create_user(client, admin)
    user = created["user"]
    member = _login(client, user["username"], created["initial_password"])
    assert client.post("/api/auth/password", headers=member,
                       json={"current_password": created["initial_password"],
                             "new_password": "member-password"}).status_code == 200

    response = client.post("/api/me/inquiries", headers=member,
                           json={"title": "근무 문의", "body": "확인 부탁드립니다."})
    assert response.status_code == 201
    inquiry = response.json()
    assert inquiry["kind"] == "inquiry" and inquiry["user_id"] == user["id"]
    assert inquiry["replies"] == []
    user_reply = client.post(f"/api/me/inquiries/{inquiry['id']}/replies", headers=member,
                             json={"body": "추가 설명입니다."})
    assert user_reply.status_code == 200
    assert user_reply.json()["replies"][0]["from_admin"] is False

    admin_list = client.get("/api/admin/inquiries?kind=inquiry", headers=admin)
    assert admin_list.status_code == 200
    assert admin_list.json()[0]["user"] == {
        "id": user["id"], "username": user["username"], "display_name": "직원"
    }
    admin_reply = client.post(f"/api/admin/inquiries/{inquiry['id']}/replies", headers=admin,
                              json={"body": "답변드립니다."})
    assert admin_reply.status_code == 200
    assert admin_reply.json()["replies"][-1]["from_admin"] is True

    from app.main import chat_feedback
    from app.schemas import ChatFeedbackRequest

    chat_feedback(ChatFeedbackRequest(rating="down", reason="근거 확인 필요", question="질문",
                                      answer="답변", evidence=[{"filename": "규정.pdf"}]), db,
                  db.get(User, user["id"]))
    feedback = client.get("/api/admin/inquiries?kind=feedback", headers=admin)
    assert feedback.status_code == 200 and len(feedback.json()) == 1
    assert feedback.json()[0]["title"] == "답변 피드백"
    assert feedback.json()[0]["body"] == "근거 확인 필요"


def test_feedback_backfill_is_idempotent(api):
    _, db, _ = api
    user = User(username="001", password_hash="hash", display_name="직원")
    db.add(user)
    db.flush()
    event = AuditEvent(actor=user.id, action="chat.feedback", entity_type="chat_answer",
                       entity_id="answer-1", payload_json={"rating": "up", "reason": None,
                       "question": "q", "answer": "a", "evidence": []})
    db.add(event)
    db.commit()

    management.backfill_feedback_inquiries(db)
    management.backfill_feedback_inquiries(db)
    inquiries = db.scalars(select(Inquiry).where(Inquiry.source_audit_id == event.id)).all()
    assert len(inquiries) == 1
    assert inquiries[0].body == "" and inquiries[0].rating == "up"
