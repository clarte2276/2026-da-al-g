"""6호선 수색승무사업소 근무계획/행로표 엑셀 파서."""
from __future__ import annotations

import io
import json
import re
from calendar import monthrange
from copy import deepcopy
from datetime import date, datetime, time, timedelta
from functools import lru_cache
from pathlib import Path

from sqlalchemy import select
from sqlalchemy.orm import Session

from .models import DutyOverride, DutyRoute, DutySchedule

OFFICE, LINE = "수색승무사업소", "6호선"
ROUTE_SHEETS = {"평일, 평평240801": "평평", "평휴240801": "평휴", "휴평240801": "휴평", "휴일,휴휴": "휴휴"}
# 수색 grid columns (A..N); 신내 block starts at O and is ignored.
STATIONS = {3: "응암", 5: "DMC", 6: "DMC", 7: "공덕", 8: "공덕", 9: "안암", 10: "안암",
            11: "봉화산", 12: "봉화산", 13: "신내기지"}
STATION_ORDER = ["응암", "DMC", "공덕", "안암", "봉화산", "신내기지"]
TIME_RE = re.compile(r"^(\d{1,2}):(\d{2})(?::(\d{2}))?$")


def _secs(value) -> int | None:
    if isinstance(value, timedelta):
        return int(value.total_seconds())
    if isinstance(value, datetime):
        value = value.time()
    if isinstance(value, time):
        return value.hour * 3600 + value.minute * 60 + value.second
    if isinstance(value, str) and (m := TIME_RE.match(value.strip())):
        return int(m[1]) * 3600 + int(m[2]) * 60 + int(m[3] or 0)
    return None


def _hhmm(secs: int) -> str:
    return f"{secs % 86400 // 3600:02d}:{secs % 3600 // 60:02d}"


def _label(value) -> str:
    return re.sub(r"\s+", "", value) if isinstance(value, str) else ""


def _block(rows: list[tuple], start: int, end: int) -> dict | None:
    """Parse one dia block: rows[start] has '출근시각', rows[end] the '근무시간' header."""
    col_a = [rows[i][0] for i in range(start + 1, end)]
    dia = next((m[1] for v in col_a if (m := re.match(r"(\d+)(\s|$)", str(v)))
                and not isinstance(v, (time, timedelta))), None)
    depart = next((s for v in col_a if (s := _secs(v)) is not None), None)
    if dia is None or depart is None or end + 3 > len(rows):
        return None
    ends, starts, totals = rows[end + 1], rows[end + 2], rows[end + 3]
    if not _secs(totals[1]):  # empty placeholder dia
        return None
    arrive = next((s for v in (ends[0], starts[0]) if (s := _secs(v)) is not None), None)
    if arrive is None:
        return None
    if arrive < depart:
        arrive += 86400

    def norm(secs: int) -> int:  # place a clock time on the turn's timeline
        while secs < depart - 3600:
            secs += 86400
        return secs

    cells: dict[int, list[tuple[int, int]]] = {}
    trains: list[tuple[int, int, str]] = []
    for r in range(start, end - 1):
        for c, v in enumerate(rows[r][:14], 1):
            if c == 1:
                continue
            if isinstance(v, int) and 6000 <= v <= 6999:
                trains.append((r, c, str(v)))
            elif (s := _secs(v)) is not None and c in STATIONS:
                cells.setdefault(s % 86400, []).append((r, c))

    def where(secs: int) -> tuple[str | None, int | None, int | None]:
        hit = cells.get(secs % 86400)
        return (STATIONS[hit[0][1]], *hit[0]) if hit else (None, None, None)

    segments = []
    for kind, cols in (("run", (3, 4, 5)), ("standby", (7, 8))):
        for c in cols:
            a, b = _secs(starts[c]), _secs(ends[c])
            if a is None or b is None or a == b:
                continue
            a, b = norm(a), norm(b)
            if b < a:
                b += 86400
            src, row, col = where(a)
            dst = where(b)[0]
            train = None
            if kind == "run" and row is not None:
                near = [t for t in trains if abs(t[0] - row) <= 2]
                train = min(near, key=lambda t: (abs(t[0] - row) + (t[0] > row), abs(t[1] - col)))[2] if near else None
            direction = None
            if kind == "run" and src and dst and src != dst:
                direction = "up" if STATION_ORDER.index(dst) > STATION_ORDER.index(src) else "down"
            segments.append({"kind": kind, "train": train, "from": src, "to": dst if kind == "run" else src,
                             "dep": _hhmm(a), "arr": _hhmm(b), "dir": direction,
                             "nextDay": a >= 86400, "_start": a})
    segments.sort(key=lambda s: s["_start"])
    for i, s in enumerate(segments):
        del s["_start"]
        if s["kind"] == "standby" and not s["from"] and i:  # wait where the previous run ended
            s["from"] = s["to"] = segments[i - 1]["to"]
    runs = [s for s in segments if s["kind"] == "run"]
    minutes = lambda v: (_secs(v) or 0) // 60
    stations: list[str] = []
    for s in runs:
        for st in (s["from"], s["to"]):
            if st and (not stations or stations[-1] != st):
                stations.append(st)
    waits = [s for s in segments if s["kind"] == "standby"]
    return {"id": dia, "depart": _hhmm(depart), "arrive": _hhmm(arrive), "nextDay": arrive >= 86400,
            "firstRun": runs[0]["dep"] if runs else "",
            "break": ", ".join(f"{s['dep']}~{s['arr']}" for s in waits),
            "route": "→".join(stations), "trainNo": (runs[0]["train"] or "") if runs else "",
            "workMinutes": minutes(totals[1]), "driveMinutes": minutes(totals[2]),
            "standbyMinutes": minutes(totals[6]), "segments": segments}


def parse_route_xlsx(data: bytes, filename: str) -> dict:
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError("행로표 엑셀 파일을 열 수 없습니다") from exc
    try:
        turns: dict[str, dict] = {}
        for sheet_name, kind in ROUTE_SHEETS.items():
            if sheet_name not in workbook.sheetnames:
                continue
            rows = [tuple(r) + (None,) * (14 - len(r)) for r in
                    workbook[sheet_name].iter_rows(max_col=14, values_only=True)]
            starts = [i for i, r in enumerate(rows) if _label(r[0]) == "출근시각"]
            for i, start in enumerate(starts):
                stop = starts[i + 1] if i + 1 < len(starts) else len(rows)
                end = next((j for j in range(start, stop) if _label(rows[j][1]) == "근무시간"), None)
                turn = _block(rows, start, end) if end is not None else None
                if turn:
                    turn["dayType"] = kind
                    turn["type"] = "night" if int(turn["id"]) >= 60 else "day"
                    turns.setdefault(f"{kind}:{turn['id']}", turn)
        if not turns:
            raise ValueError("행로표 시트(평일, 평평240801 등)를 찾을 수 없습니다")
        found = re.search(r"(\d{2})(\d{2})(\d{2})", filename)
        effective = f"20{found[1]}-{found[2]}-{found[3]}" if found else datetime.now().astimezone().date().isoformat()
        return {"effective": effective, "source": filename, "turns": turns}
    finally:
        workbook.close()


def classify(code: str) -> str:
    if code == "비" or re.fullmatch(r"\d+~", code) or re.fullmatch(r"대\d+~", code):
        return "off"
    if re.fullmatch(r"(휴|운휴|대휴)\d+~?", code):
        return "rest"
    if re.fullmatch(r"대\d+", code):
        return "standby"
    if code.startswith("지정") or code in {"주간", "보완교육"}:
        return "designated"
    if code.isdigit():
        return "day" if int(code) < 60 else "night"
    return "unknown"


def _day_types(workbook, first: date, last: date) -> dict[str, str]:
    holidays = {date.fromisoformat(value) for value in _holiday_dates(workbook)}
    out, day = {}, first
    while day <= last:
        out[day.isoformat()] = "휴" if day.weekday() >= 5 or day in holidays else "평"
        day += timedelta(days=1)
    return out


def _holiday_dates(workbook) -> list[str]:
    if "공휴일" not in workbook.sheetnames:
        return []
    return sorted({(value.date() if isinstance(value, datetime) else value).isoformat()
                   for (value,) in workbook["공휴일"].iter_rows(max_col=1, values_only=True)
                   if isinstance(value, (date, datetime))})


def parse_duty_xlsx(data: bytes, month: str, filename: str) -> dict:
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        if "교번근무계획(지정근무반영)" not in workbook.sheetnames:
            workbook["교번기관사 근무계획"]
    except Exception as exc:
        raise ValueError("근무계획 엑셀 양식이 올바르지 않습니다") from exc
    try:
        year, mon = int(month[:4]), int(month[-2:])
        days_in = monthrange(year, mon)[1]
        drivers = (_designated_drivers(workbook["교번근무계획(지정근무반영)"], year, mon, days_in)
                   if "교번근무계획(지정근무반영)" in workbook.sheetnames
                   else _legacy_drivers(workbook["교번기관사 근무계획"], month, days_in))
        if not drivers or len({x["no"] for x in drivers}) != len(drivers):
            raise ValueError("기관사 행이 없거나 연번이 중복되었습니다")
        first = date(year, mon, 1)
        return {"month": month, "line": LINE, "office": OFFICE, "source": filename,
                "dates": [(first + timedelta(days=i)).isoformat() for i in range(days_in)],
                "dayTypes": _day_types(workbook, first - timedelta(days=1), first + timedelta(days=days_in)),
                "holidayDates": _holiday_dates(workbook),
                "drivers": drivers, "turns": {}, "turnDetailIsDummy": True}
    finally:
        workbook.close()


def _day(code: str) -> dict:
    kind = classify(code)
    return {"code": code, "type": kind, "turn": code if kind in {"day", "night", "standby"} else None}


def _designated_drivers(sheet, year: int, mon: int, days_in: int) -> list[dict]:
    rows = sheet.iter_rows(values_only=True)
    header = next((r for r in rows if "승무원" in r), None)
    if header is None:
        raise ValueError("교번근무계획 시트의 머리글을 찾을 수 없습니다")
    name_col = header.index("승무원")
    emp_col = next((i for i, v in enumerate(header) if isinstance(v, str) and v.endswith("사원번호")), None)
    day_cols = {int(v[:-1]): i for i, v in enumerate(header)
                if isinstance(v, str) and re.fullmatch(r"\d{1,2}일", v) and int(v[:-1]) <= days_in}
    if len(day_cols) != days_in:
        raise ValueError("선택한 월의 날짜 열을 찾을 수 없습니다")
    drivers = []
    for row in rows:
        no, name = row[0], row[name_col]
        if not isinstance(no, int) or not isinstance(name, str) or not name.strip():
            continue
        days = {}
        for d, col in day_cols.items():
            code = str(row[col]).strip() if row[col] is not None else ""  # blank = unknown (e.g. 휴직)
            days[date(year, mon, d).isoformat()] = _day(code)
        emp = row[emp_col] if emp_col is not None else None
        drivers.append({"no": no, "name": name.strip(), "empNo": str(emp) if emp is not None else None,
                        "days": days})
    return drivers


def _legacy_drivers(sheet, month: str, days_in: int) -> list[dict]:
    header = next(sheet.iter_rows(min_row=2, max_row=2, values_only=True))
    dates = {i: v.date() if isinstance(v, datetime) else v for i, v in enumerate(header)
             if isinstance(v, (date, datetime)) and v.strftime("%Y-%m") == month}
    if len(dates) != days_in:
        raise ValueError("선택한 월의 날짜 열을 찾을 수 없습니다")
    drivers = []
    for row in sheet.iter_rows(min_row=7, values_only=True):
        no, name = row[:2]
        if not isinstance(no, int) or not isinstance(name, str) or not name.strip():
            continue
        days = {}
        for col, day in dates.items():
            code = str(row[col]).strip() if row[col] is not None else ""
            if not code:
                raise ValueError(f"{name}: {day} 근무 코드가 비어 있습니다")
            days[day.isoformat()] = _day(code)
        drivers.append({"no": no, "name": name.strip(), "days": days})
    return drivers


def merge_routes(schedule: dict, routes: dict | None) -> dict:
    """Attach dayType/turnKey/continues per day and the referenced 행로표 turns."""
    out = {**schedule, "drivers": [], "turns": deepcopy(schedule.get("turns") or {})}
    types = schedule.get("dayTypes") or {}

    def dtype(day: date) -> str:
        return types.get(day.isoformat()) or ("휴" if day.weekday() >= 5 else "평")

    table = (routes or {}).get("turns", {})
    for driver in schedule["drivers"]:
        days = {}
        for iso, cell in driver["days"].items():
            day, code = date.fromisoformat(iso), cell["code"]
            single, combo = dtype(day), dtype(day) + dtype(day + timedelta(days=1))
            number = re.fullmatch(r"대?(\d+)", code)
            night = cell["type"] == "night" or (cell["type"] == "standby" and number and int(number[1]) >= 60)
            key = None
            if cell["type"] == "day":
                key = f"{single * 2}:{code}"
            elif cell["type"] == "night":
                key = f"{combo}:{code}"
            extra = {"dayType": combo if night else single, "turnKey": key}
            if (m := re.fullmatch(r"(\d+)~", code)) and int(m[1]) >= 60:
                extra["continues"] = f"{dtype(day - timedelta(days=1))}{single}:{m[1]}"
            for k in (key, extra.get("continues")):
                if k and k in table:
                    out["turns"][k] = table[k]
            days[iso] = {**cell, **extra}
        out["drivers"].append({**driver, "days": days})
    if routes:
        out["turnDetailIsDummy"] = False
    pack = dia_data()
    for key, turn in out["turns"].items():
        table_key = table_key_for_turn(key)
        code = key.rsplit(":", 1)[-1]
        detail = pack["schedules"].get(table_key, {}).get(code, {}) if table_key else {}
        note = detail.get("note") or ""
        last = 0
        turn["trainTimes"] = []
        for match in re.finditer(r"(\d{1,2}):(\d{2})", note):
            turn["trainTimes"].append({"time": f"{int(match[1]):02d}:{match[2]}",
                                       "depot": "출" in note[last:match.start()]})
            last = match.end()
        turn["diaImage"] = f"/api/duty/dia-images/{table_key}-{code}.png" if table_key and re.fullmatch(r"\d{1,2}", code) else None
    return out


@lru_cache(maxsize=1)
def dia_data() -> dict:
    return json.loads((Path(__file__).parent / "data" / "dia-data.json").read_text(encoding="utf-8"))


def table_key_for_turn(turn_key: str | None) -> str | None:
    if not turn_key or ":" not in turn_key:
        return None
    kind, code = turn_key.split(":", 1)
    try:
        night = int(code) >= 60
    except ValueError:
        return None
    return {"평평": "ordord" if night else "ord", "휴휴": "holhol" if night else "hol",
            "평휴": "ordhol", "휴평": "holord"}.get(kind)


def _month_after(month: str, count: int) -> str:
    year, mon = map(int, month.split("-"))
    year, mon = divmod(year * 12 + mon - 1 + count, 12)
    return f"{year:04d}-{mon + 1:02d}"


def projected_months(latest_month: str | None) -> list[str]:
    return [_month_after(latest_month, i) for i in range(1, 7)] if latest_month else []


def _cycle_code(pack: dict, index: int) -> str:
    cycle = pack["cycle"]
    entry = cycle[index % len(cycle)]["code"]
    if entry != "~":
        return entry
    previous = cycle[(index - 1) % len(cycle)]["code"]
    return f"{previous}~" if previous.isdigit() and int(previous) >= 60 else "비"


def _known_holidays(pack: dict, schedules: list[DutySchedule]) -> tuple[set[str], set[int]]:
    holidays = set(pack.get("holidays", []))
    for row in schedules:
        data = row.data
        holidays.update(data.get("holidayDates", []))
        for iso, kind in data.get("dayTypes", {}).items():
            if kind == "휴" and date.fromisoformat(iso).weekday() < 5:
                holidays.add(iso)
    return holidays, {int(day[:4]) for day in holidays}


def project_schedule(latest: dict, month: str, uploads: list[DutySchedule] = ()) -> dict:
    """Project one month by aligning each latest-upload roster to the 94-day cycle."""
    pack = dia_data()
    length = pack["metadata"]["cycleLength"]
    anchor = date.fromisoformat(f"{latest['month']}-01")
    year, mon = map(int, month.split("-"))
    first = date(year, mon, 1)
    days_in = monthrange(year, mon)[1]
    holidays, known_years = _known_holidays(pack, list(uploads))
    dates = [date(year, mon, day) for day in range(1, days_in + 1)]
    drivers = []
    for driver in latest["drivers"]:
        samples = [(date.fromisoformat(iso), cell["code"]) for iso, cell in driver.get("days", {}).items()
                   if cell.get("code") and classify(cell["code"]) in {"day", "night", "off", "rest", "standby"}]
        if not samples:
            continue
        best_index, best_score = max(
            ((start, sum(_cycle_code(pack, start + (day - anchor).days) == code for day, code in samples) / len(samples))
             for start in range(length)), key=lambda result: result[1])
        if best_score < 0.6:
            continue
        days = {}
        for day in dates:
            iso = day.isoformat()
            code = _cycle_code(pack, best_index + (day - anchor).days)
            days[iso] = _day(code)
        drivers.append({key: value for key, value in driver.items() if key != "days"} | {"days": days})
    day_types = {}
    day = first - timedelta(days=1)
    last = dates[-1] + timedelta(days=1)
    while day <= last:
        if day.weekday() >= 5:
            kind = "휴"
        else:
            # ponytail: missing holiday years use weekends only, add that year's workbook dates when available.
            kind = "휴" if day.year in known_years and day.isoformat() in holidays else "평"
        day_types[day.isoformat()] = kind
        day += timedelta(days=1)
    return {"month": month, "line": LINE, "office": OFFICE, "source": "94-day cycle",
            "dates": [day.isoformat() for day in dates], "dayTypes": day_types,
            "drivers": drivers, "turns": {}, "turnDetailIsDummy": True}


def load_duty(db: Session, month: str) -> dict | None:
    row = db.get(DutySchedule, month)
    projected, projected_from = False, None
    if row:
        schedule = row.data
    else:
        uploads = list(db.scalars(select(DutySchedule).order_by(DutySchedule.month.desc())).all())
        if not uploads or month not in projected_months(uploads[0].month):
            return None
        projected, projected_from = True, uploads[0].month
        schedule = project_schedule(uploads[0].data, month, uploads)
    route = db.scalars(select(DutyRoute).order_by(DutyRoute.effective_date.desc()).limit(1)).first()
    start = (date.fromisoformat(schedule["dates"][0]) - timedelta(days=1)).isoformat()
    end = schedule["dates"][-1]
    overrides = list(db.scalars(select(DutyOverride).where(
        DutyOverride.active.is_(True), DutyOverride.date >= start, DutyOverride.date <= end)).all())
    result = _apply_overrides(schedule, overrides, route.data if route else None)
    result["projected"] = projected
    result["projectedFrom"] = projected_from
    return result


def _apply_overrides(schedule: dict, overrides: list[DutyOverride], routes: dict | None) -> dict:
    if not overrides:
        return merge_routes(schedule, routes)
    updated = deepcopy(schedule)
    by_day = {(row.driver_name, row.date): row for row in overrides}
    drivers = {driver["name"]: driver for driver in updated["drivers"]}
    original_codes = {}
    for row in overrides:
        driver = drivers.get(row.driver_name)
        cell = driver and driver["days"].get(row.date)
        if not cell:
            continue
        original_codes[(row.driver_name, row.date)] = row.original_code
        cell.update(_day(row.code))
    for (name, iso), original_code in original_codes.items():
        driver = drivers[name]
        next_iso = (date.fromisoformat(iso) + timedelta(days=1)).isoformat()
        next_cell = driver["days"].get(next_iso)
        row = by_day[(name, iso)]
        if not next_cell or (name, next_iso) in by_day:
            continue
        if (classify(original_code) == "night" or classify(row.code) == "night") \
                and (next_cell["code"] == f"{original_code}~" or next_cell["code"] == "비"):
            next_cell.update(_day(f"{row.code}~" if classify(row.code) == "night" else "비"))
    result = merge_routes(updated, routes)
    for row in overrides:
        driver = next((item for item in result["drivers"] if item["name"] == row.driver_name), None)
        if not driver or row.date not in driver["days"]:
            continue
        before = deepcopy(updated)
        before_driver = next(item for item in before["drivers"] if item["name"] == row.driver_name)
        before_driver["days"][row.date].update(_day(row.original_code))
        original_day = merge_routes(before, routes)["drivers"][
            next(i for i, item in enumerate(before["drivers"]) if item["name"] == row.driver_name)
        ]["days"][row.date]
        driver["days"][row.date]["original"] = {
            key: original_day[key] for key in ("code", "type", "dayType", "turnKey", "continues") if key in original_day
        }
        driver["days"][row.date]["swapPostId"] = row.post_id
    return result
