from __future__ import annotations

import hashlib
import re
import secrets
from datetime import UTC, datetime
from io import BytesIO
from typing import Annotated, Any, Literal

from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile
from fastapi.responses import Response
from openpyxl import Workbook, load_workbook
from pydantic import BaseModel, Field
from sqlalchemy import select, update
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from .auth import AdminUser, CurrentUser, normalize_username, password_hash
from .config import get_settings
from .db import get_db
from .duty import merge_routes, parse_duty_xlsx, parse_route_xlsx
from .models import (
    AuditEvent,
    AuthSession,
    BookmarkRecord,
    ConversationRecord,
    DutyRoute,
    DutySchedule,
    Inquiry,
    InquiryReply,
    User,
)
from .schemas import AdminInquiryOut, InquiryOut, UserCredentialsOut, UserOut

router = APIRouter(prefix="/api")
DB = Annotated[Session, Depends(get_db)]
MONTH_RE = re.compile(r"^\d{4}-(0[1-9]|1[0-2])$")
MEMBER_ID_RE = re.compile(r"^\d+$")
PASSWORD_ALPHABET = "ABCDEFGHJKLMNPQRSTUVWXYZabcdefghijkmnopqrstuvwxyz23456789"
XLSX_MIME_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


class ConversationWrite(BaseModel):
    id: str = Field(min_length=1, max_length=100)
    title: str = Field(max_length=200)
    createdAt: datetime
    messages: list[dict[str, Any]] = Field(max_length=500)


class BookmarkWrite(BaseModel):
    regulation: str = Field(min_length=1, max_length=200)
    chapter: str = Field(min_length=1, max_length=200)
    version: str = Field(max_length=100)
    excerpt: str | None = Field(default=None, max_length=10000)
    content: str | None = Field(default=None, max_length=20000)
    sourcePath: str | None = Field(default=None, max_length=2000)
    chunkId: str | None = Field(default=None, max_length=100)
    retriever: str | None = Field(default=None, max_length=100)
    documentId: str | None = Field(default=None, max_length=100)
    versionId: str | None = Field(default=None, max_length=100)
    page: int | None = Field(default=None, ge=1)


class ImportWrite(BaseModel):
    conversations: list[ConversationWrite] = Field(default_factory=list, max_length=1000)
    bookmarks: list[BookmarkWrite] = Field(default_factory=list, max_length=1000)


class UserCreate(BaseModel):
    username: str = Field(min_length=1)
    display_name: str = Field(min_length=1, max_length=100)


class UserUpdate(BaseModel):
    display_name: str | None = Field(default=None, min_length=1, max_length=100)
    is_active: bool | None = None
    reset_password: bool = False


class CredentialRow(BaseModel):
    username: str
    display_name: str
    initial_password: str


class CredentialsRequest(BaseModel):
    rows: list[CredentialRow]


class InquiryCreate(BaseModel):
    title: str = Field(min_length=1, max_length=200)
    body: str = Field(min_length=1, max_length=4000)


class InquiryReplyCreate(BaseModel):
    body: str = Field(min_length=1, max_length=4000)


def _audit(db: Session, actor: str, action: str, entity_type: str, entity_id: str,
           payload: dict | None = None) -> None:
    db.add(AuditEvent(actor=actor, action=action, entity_type=entity_type,
                      entity_id=entity_id, payload_json=payload or {}))


def _bookmark_id(data: BookmarkWrite) -> str:
    return hashlib.sha256(f"{data.regulation}|{data.chapter}".encode()).hexdigest()


def _validate_conversation(data: ConversationWrite) -> None:
    for message in data.messages:
        if not isinstance(message.get("isUser"), bool) or not isinstance(message.get("text"), str):
            raise HTTPException(422, "Invalid conversation message")
        if len(message["text"]) > 20000:
            raise HTTPException(422, "Conversation message too long")


@router.get("/me/conversations")
def my_conversations(db: DB, user: CurrentUser) -> list[dict]:
    rows = db.scalars(select(ConversationRecord).where(ConversationRecord.user_id == user.id,
                            ConversationRecord.deleted.is_(False)).order_by(ConversationRecord.updated_at.desc())).all()
    return [row.data for row in rows]


@router.put("/me/conversations/{conversation_id}")
def save_conversation(conversation_id: str, payload: ConversationWrite, db: DB, user: CurrentUser) -> dict:
    if conversation_id != payload.id:
        raise HTTPException(422, "Conversation ID mismatch")
    _validate_conversation(payload)
    row = db.get(ConversationRecord, (conversation_id, user.id))
    if row and row.deleted:
        raise HTTPException(410, "Conversation was deleted")
    if not row:
        row = ConversationRecord(id=conversation_id, user_id=user.id, data={})
        db.add(row)
    row.data = payload.model_dump(mode="json")
    db.commit()
    return row.data


@router.delete("/me/conversations/{conversation_id}")
def delete_conversation(conversation_id: str, db: DB, user: CurrentUser) -> dict:
    row = db.get(ConversationRecord, (conversation_id, user.id))
    if not row:
        row = ConversationRecord(id=conversation_id, user_id=user.id, data={}, deleted=True)
        db.add(row)
    row.deleted = True
    row.data = {}
    db.commit()
    return {"status": "ok"}


@router.get("/me/bookmarks")
def my_bookmarks(db: DB, user: CurrentUser) -> list[dict]:
    rows = db.scalars(select(BookmarkRecord).where(BookmarkRecord.user_id == user.id,
                            BookmarkRecord.deleted.is_(False)).order_by(BookmarkRecord.updated_at.desc())).all()
    return [row.data for row in rows]


@router.post("/me/bookmarks")
def save_bookmark(payload: BookmarkWrite, db: DB, user: CurrentUser) -> dict:
    key = _bookmark_id(payload)
    row = db.get(BookmarkRecord, (key, user.id))
    if row and row.deleted:
        raise HTTPException(410, "Bookmark was deleted")
    if not row:
        row = BookmarkRecord(id=key, user_id=user.id, data={})
        db.add(row)
    row.data = payload.model_dump(mode="json")
    db.commit()
    return row.data


@router.delete("/me/bookmarks")
def delete_bookmark(payload: BookmarkWrite, db: DB, user: CurrentUser) -> dict:
    key = _bookmark_id(payload)
    row = db.get(BookmarkRecord, (key, user.id))
    if not row:
        row = BookmarkRecord(id=key, user_id=user.id, data={}, deleted=True)
        db.add(row)
    row.deleted = True
    row.data = {}
    db.commit()
    return {"status": "ok"}


@router.post("/me/import")
def import_local(payload: ImportWrite, db: DB, user: CurrentUser) -> dict:
    for item in payload.conversations:
        _validate_conversation(item)
        if not db.get(ConversationRecord, (item.id, user.id)):
            db.add(ConversationRecord(id=item.id, user_id=user.id, data=item.model_dump(mode="json")))
    for item in payload.bookmarks:
        key = _bookmark_id(item)
        if not db.get(BookmarkRecord, (key, user.id)):
            db.add(BookmarkRecord(id=key, user_id=user.id, data=item.model_dump(mode="json")))
    db.commit()
    return {"status": "ok"}


@router.get("/duty/{month}")
def get_duty(month: str, db: DB, _user: CurrentUser) -> dict:
    if not MONTH_RE.fullmatch(month):
        raise HTTPException(422, "Invalid month")
    row = db.get(DutySchedule, month)
    if not row:
        raise HTTPException(404, "No duty schedule for this month")
    route = db.scalars(select(DutyRoute).order_by(DutyRoute.effective_date.desc()).limit(1)).first()
    return merge_routes(row.data, route.data if route else None)


@router.get("/duty")
def available_duty_months(db: DB, _user: CurrentUser) -> list[str]:
    return list(db.scalars(select(DutySchedule.month).order_by(DutySchedule.month.desc())).all())


@router.get("/admin/duty")
def list_duty(db: DB, _admin: AdminUser) -> list[dict]:
    rows = db.scalars(select(DutySchedule).order_by(DutySchedule.month.desc())).all()
    return [{"month": x.month, "source_filename": x.source_filename, "updated_at": x.updated_at,
             "driver_count": len(x.data["drivers"])} for x in rows]


@router.post("/admin/duty/{month}")
def upload_duty(month: str, file: Annotated[UploadFile, File(...)], db: DB, admin: AdminUser) -> dict:
    if not MONTH_RE.fullmatch(month) or not (file.filename or "").lower().endswith(".xlsx"):
        raise HTTPException(422, "YYYY-MM month and .xlsx file required")
    data = file.file.read(get_settings().max_upload_bytes + 1)
    if len(data) > get_settings().max_upload_bytes:
        raise HTTPException(413, "File too large")
    try:
        parsed = parse_duty_xlsx(data, month, file.filename or "schedule.xlsx")
    except ValueError as exc:
        raise HTTPException(422, str(exc)) from exc
    row = db.get(DutySchedule, month)
    if not row:
        row = DutySchedule(month=month, data=parsed, source_filename=file.filename or "schedule.xlsx")
        db.add(row)
    else:
        row.data = parsed
        row.source_filename = file.filename or "schedule.xlsx"
        row.updated_at = datetime.now(UTC)
    _audit(db, admin.id, "duty.upload", "duty_schedule", month)
    db.commit()
    return {"month": month, "driver_count": len(parsed["drivers"])}


@router.get("/admin/duty-routes")
def list_duty_routes(db: DB, _admin: AdminUser) -> list[dict]:
    rows = db.scalars(select(DutyRoute).order_by(DutyRoute.effective_date.desc())).all()
    return [{"id": x.id, "effective_date": x.effective_date, "source_filename": x.source_filename,
             "updated_at": x.updated_at, "turn_count": len(x.data["turns"])} for x in rows]


def save_duty_route(db: Session, parsed: dict, filename: str) -> DutyRoute:
    row = db.scalar(select(DutyRoute).where(DutyRoute.effective_date == parsed["effective"]))
    if not row:
        row = DutyRoute(effective_date=parsed["effective"], data=parsed, source_filename=filename)
        db.add(row)
    row.data, row.source_filename, row.updated_at = parsed, filename, datetime.now(UTC)
    return row


@router.post("/admin/duty-routes")
def upload_duty_route(file: Annotated[UploadFile, File(...)], db: DB, admin: AdminUser) -> dict:
    if not (file.filename or "").lower().endswith(".xlsx"):
        raise HTTPException(422, ".xlsx file required")
    data = file.file.read(get_settings().max_upload_bytes + 1)
    if len(data) > get_settings().max_upload_bytes:
        raise HTTPException(413, "File too large")
    try:
        parsed = parse_route_xlsx(data, file.filename or "routes.xlsx")
    except ValueError as exc:
        raise HTTPException(422, str(exc)) from exc
    save_duty_route(db, parsed, file.filename or "routes.xlsx")
    _audit(db, admin.id, "duty_route.upload", "duty_route", parsed["effective"])
    db.commit()
    return {"effective_date": parsed["effective"], "turn_count": len(parsed["turns"])}


@router.get("/admin/users", response_model=list[UserOut])
def list_users(db: DB, _admin: AdminUser) -> list[User]:
    return list(db.scalars(select(User).where(User.role == "user").order_by(User.created_at.desc())).all())


def _temporary_password() -> str:
    return "".join(secrets.choice(PASSWORD_ALPHABET) for _ in range(14))


def _make_user(username: str, display_name: str) -> tuple[User, str]:
    initial_password = _temporary_password()
    return User(
        username=username,
        display_name=display_name,
        password_hash=password_hash.hash(initial_password),
        role="user",
        is_active=True,
        must_change_password=True,
    ), initial_password


def _xlsx_response(workbook: Workbook, filename: str) -> Response:
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return Response(
        output.getvalue(),
        media_type=XLSX_MIME_TYPE,
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.get("/admin/users/template.xlsx")
def user_template(_admin: AdminUser) -> Response:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["사번", "이름"])
    return _xlsx_response(workbook, "users-template.xlsx")


@router.post("/admin/users", response_model=UserCredentialsOut, response_model_exclude_none=True, status_code=201)
def create_user(payload: UserCreate, db: DB, admin: AdminUser) -> dict:
    username = normalize_username(payload.username)
    if not MEMBER_ID_RE.fullmatch(username):
        raise HTTPException(422, "사번은 숫자만 입력해야 합니다.")
    if db.scalar(select(User).where(User.username == username)):
        raise HTTPException(409, "이미 존재하는 사번입니다.")
    if not payload.display_name.strip():
        raise HTTPException(422, "이름을 입력해야 합니다.")
    user, initial_password = _make_user(username, payload.display_name.strip())
    db.add(user)
    try:
        db.flush()
    except IntegrityError as exc:
        db.rollback()
        raise HTTPException(409, "이미 존재하는 사번입니다.") from exc
    _audit(db, admin.id, "user.created", "user", user.id)
    db.commit()
    return {"user": user, "initial_password": initial_password}


@router.patch("/admin/users/{user_id}", response_model=UserCredentialsOut, response_model_exclude_none=True)
def update_user(user_id: str, payload: UserUpdate, db: DB, admin: AdminUser) -> dict:
    user = db.get(User, user_id)
    if not user or user.role != "user":
        raise HTTPException(404, "User not found")
    if payload.display_name is not None:
        if not payload.display_name.strip():
            raise HTTPException(422, "Display name is required")
        user.display_name = payload.display_name.strip()
    if payload.is_active is not None:
        user.is_active = payload.is_active
    initial_password = None
    if payload.reset_password:
        initial_password = _temporary_password()
        user.password_hash = password_hash.hash(initial_password)
        user.must_change_password = True
    if payload.reset_password or payload.is_active is False:
        db.execute(update(AuthSession).where(AuthSession.user_id == user.id,
                   AuthSession.revoked_at.is_(None)).values(revoked_at=datetime.now(UTC)))
    _audit(db, admin.id, "user.updated", "user", user.id)
    db.commit()
    return {"user": user, "initial_password": initial_password}


@router.post("/admin/users/import")
def import_users(file: Annotated[UploadFile, File(...)], db: DB, admin: AdminUser) -> dict:
    if not (file.filename or "").lower().endswith(".xlsx"):
        raise HTTPException(422, ".xlsx 파일이 필요합니다.")
    data = file.file.read(get_settings().max_upload_bytes + 1)
    if len(data) > get_settings().max_upload_bytes:
        raise HTTPException(413, "파일이 너무 큽니다.")
    try:
        workbook = load_workbook(BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        raise HTTPException(422, "올바른 .xlsx 파일이 아닙니다.") from exc

    created, skipped = [], []
    try:
        sheet = workbook.active
        header = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
        if tuple(header[:2]) != ("사번", "이름"):
            raise HTTPException(422, "첫 행에 사번과 이름 열이 필요합니다.")
        seen: set[str] = set()
        for values in sheet.iter_rows(min_row=2, values_only=True):
            raw_username = values[0] if values else None
            raw_name = values[1] if len(values) > 1 else None
            if raw_username is None and raw_name is None:
                continue
            if isinstance(raw_username, float) and raw_username.is_integer():
                raw_username = int(raw_username)
            username = str(raw_username).strip() if raw_username is not None else ""
            display_name = str(raw_name).strip() if raw_name is not None else ""
            if not MEMBER_ID_RE.fullmatch(username):
                skipped.append({"username": username, "reason": "사번은 숫자만 입력해야 합니다."})
                continue
            if not display_name:
                skipped.append({"username": username, "reason": "이름을 입력해야 합니다."})
                continue
            if len(display_name) > 100:
                skipped.append({"username": username, "reason": "이름은 100자 이하여야 합니다."})
                continue
            if username in seen or db.scalar(select(User.id).where(User.username == username)):
                skipped.append({"username": username, "reason": "이미 존재하는 사번입니다."})
                continue

            user, initial_password = _make_user(username, display_name)
            try:
                with db.begin_nested():
                    db.add(user)
                    db.flush()
            except IntegrityError:
                skipped.append({"username": username, "reason": "이미 존재하는 사번입니다."})
                continue
            _audit(db, admin.id, "user.created", "user", user.id)
            seen.add(username)
            created.append({"username": username, "display_name": display_name,
                            "initial_password": initial_password})
    finally:
        workbook.close()
    db.commit()
    return {"created": created, "skipped": skipped}


@router.post("/admin/users/credentials.xlsx")
def export_credentials(payload: CredentialsRequest, _admin: AdminUser) -> Response:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["사번", "이름", "초기 비밀번호"])
    for item in payload.rows:
        sheet.append([item.username, item.display_name, item.initial_password])
        sheet.cell(sheet.max_row, 1).number_format = "@"
    return _xlsx_response(workbook, "user-credentials.xlsx")


def _inquiry_data(db: Session, inquiry: Inquiry, user: User | None = None) -> dict:
    replies = db.scalars(select(InquiryReply).where(InquiryReply.inquiry_id == inquiry.id)
                         .order_by(InquiryReply.created_at, InquiryReply.id)).all()
    result = {
        "id": inquiry.id,
        "user_id": inquiry.user_id,
        "kind": inquiry.kind,
        "title": inquiry.title,
        "body": inquiry.body,
        "rating": inquiry.rating,
        "question": inquiry.question,
        "answer": inquiry.answer,
        "evidence": inquiry.evidence,
        "created_at": inquiry.created_at,
        "replies": [{"id": reply.id, "from_admin": reply.from_admin, "body": reply.body,
                     "created_at": reply.created_at} for reply in replies],
    }
    if user:
        result["user"] = {"id": user.id, "username": user.username, "display_name": user.display_name}
    return result


@router.get("/me/inquiries", response_model=list[InquiryOut])
def my_inquiries(db: DB, user: CurrentUser) -> list[dict]:
    rows = db.scalars(select(Inquiry).where(Inquiry.user_id == user.id)
                      .order_by(Inquiry.created_at.desc())).all()
    return [_inquiry_data(db, inquiry) for inquiry in rows]


@router.post("/me/inquiries", response_model=InquiryOut, status_code=201)
def create_inquiry(payload: InquiryCreate, db: DB, user: CurrentUser) -> dict:
    inquiry = Inquiry(user_id=user.id, kind="inquiry", title=payload.title,
                      body=payload.body, evidence=[])
    db.add(inquiry)
    db.commit()
    return _inquiry_data(db, inquiry)


def _add_reply(db: Session, inquiry: Inquiry, body: str, from_admin: bool) -> dict:
    db.add(InquiryReply(inquiry_id=inquiry.id, from_admin=from_admin, body=body))
    db.commit()
    return _inquiry_data(db, inquiry)


@router.post("/me/inquiries/{inquiry_id}/replies", response_model=InquiryOut)
def reply_to_own_inquiry(inquiry_id: str, payload: InquiryReplyCreate, db: DB, user: CurrentUser) -> dict:
    inquiry = db.get(Inquiry, inquiry_id)
    if not inquiry or inquiry.user_id != user.id:
        raise HTTPException(404, "Inquiry not found")
    return _add_reply(db, inquiry, payload.body, from_admin=False)


@router.get("/admin/inquiries", response_model=list[AdminInquiryOut])
def admin_inquiries(
    db: DB,
    _admin: AdminUser,
    kind: Literal["inquiry", "feedback"] | None = Query(default=None),
) -> list[dict]:
    query = select(Inquiry, User).join(User, User.id == Inquiry.user_id)
    if kind:
        query = query.where(Inquiry.kind == kind)
    rows = db.execute(query.order_by(Inquiry.created_at.desc())).all()
    return [_inquiry_data(db, inquiry, user) for inquiry, user in rows]


@router.post("/admin/inquiries/{inquiry_id}/replies", response_model=InquiryOut)
def reply_to_inquiry(inquiry_id: str, payload: InquiryReplyCreate, db: DB, _admin: AdminUser) -> dict:
    inquiry = db.get(Inquiry, inquiry_id)
    if not inquiry:
        raise HTTPException(404, "Inquiry not found")
    return _add_reply(db, inquiry, payload.body, from_admin=True)


def backfill_feedback_inquiries(db: Session) -> None:
    rows = db.execute(
        select(AuditEvent, User)
        .join(User, User.id == AuditEvent.actor)
        .outerjoin(Inquiry, Inquiry.source_audit_id == AuditEvent.id)
        .where(AuditEvent.action == "chat.feedback", Inquiry.id.is_(None))
    ).all()
    for event, user in rows:
        payload = event.payload_json or {}
        rating = payload.get("rating")
        db.add(Inquiry(
            user_id=user.id,
            kind="feedback",
            title="답변 피드백",
            body=payload.get("reason") or "",
            rating=rating if rating in {"up", "down"} else None,
            question=payload.get("question") or "",
            answer=payload.get("answer") or "",
            evidence=payload.get("evidence") if isinstance(payload.get("evidence"), list) else [],
            source_audit_id=event.id,
            created_at=event.created_at,
        ))
    if rows:
        db.commit()
