from __future__ import annotations

from pathlib import Path
from typing import Any

from pptx import Presentation

from .base import (
    DocumentParser,
    ParsedAsset,
    ParsedDocument,
    ParsedFragment,
    clean_text,
)


def _shape_box(shape: Any) -> dict[str, Any]:
    return {
        "x_emu": int(getattr(shape, "left", 0)),
        "y_emu": int(getattr(shape, "top", 0)),
        "width_emu": int(getattr(shape, "width", 0)),
        "height_emu": int(getattr(shape, "height", 0)),
    }


def _table_payload(table: Any) -> tuple[list[list[str]], str]:
    rows: list[list[str]] = []
    for row in table.rows:
        rows.append([clean_text(cell.text) for cell in row.cells])
    html_rows = [
        "<tr>" + "".join(f"<td>{_escape_html(cell)}</td>" for cell in row) + "</tr>"
        for row in rows
    ]
    return rows, "<table>" + "".join(html_rows) + "</table>"


def _escape_html(value: str) -> str:
    return (
        value.replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
        .replace('"', "&quot;")
    )


class PptxParser(DocumentParser):
    name = "python-pptx"
    version = "1.x"

    def parse(self, path: Path) -> ParsedDocument:
        presentation = Presentation(str(path))
        fragments: list[ParsedFragment] = []
        assets: list[ParsedAsset] = []
        ordinal = 0

        for slide_number, slide in enumerate(presentation.slides, start=1):
            slide_key = f"slide:{slide_number}"
            slide_text: list[str] = []
            child_fragments: list[ParsedFragment] = []

            def visit_shape(
                shape: Any,
                shape_key: str,
                *,
                slide_text=slide_text,
                child_fragments=child_fragments,
                slide_number=slide_number,
                slide_key=slide_key,
            ) -> None:
                nonlocal ordinal
                shape_text = clean_text(getattr(shape, "text", ""))
                shape_kind = "shape"
                table_json = None
                html = None
                asset_key = None

                if getattr(shape, "has_table", False):
                    shape_kind = "table"
                    table_json, html = _table_payload(shape.table)
                    table_text = "\n".join(" | ".join(row) for row in table_json)
                    shape_text = clean_text(table_text)
                elif getattr(shape, "shape_type", None) == 13:
                    shape_kind = "image"
                    image = shape.image
                    asset_key = f"{shape_key}:image"
                    assets.append(
                        ParsedAsset(
                            key=asset_key,
                            data=image.blob,
                            extension=image.ext or "bin",
                            media_type=getattr(image, "content_type", None),
                        )
                    )

                if shape_text:
                    slide_text.append(shape_text)
                child_fragments.append(
                    ParsedFragment(
                        stable_key=shape_key,
                        kind=shape_kind,
                        ordinal=ordinal,
                        text=shape_text or None,
                        html=html,
                        table_json=table_json,
                        locator={"slide": slide_number, "shape_id": int(shape.shape_id)},
                        bbox=_shape_box(shape),
                        parent_key=slide_key,
                        asset_key=asset_key,
                        metadata={"shape_type": int(getattr(shape, "shape_type", 0))},
                    )
                )
                ordinal += 1

                nested = getattr(shape, "shapes", None)
                if nested is not None:
                    for child_index, child in enumerate(nested):
                        visit_shape(child, f"{shape_key}:child:{child_index}")

            for shape in slide.shapes:
                visit_shape(shape, f"{slide_key}:shape:{int(shape.shape_id)}")

            fragments.append(
                ParsedFragment(
                    stable_key=slide_key,
                    kind="slide",
                    ordinal=ordinal,
                    title=slide_text[0][:200] if slide_text else f"슬라이드 {slide_number}",
                    text="\n\n".join(slide_text) or None,
                    locator={"slide": slide_number},
                    metadata={"shape_count": len(slide.shapes)},
                )
            )
            ordinal += 1  # noqa: SIM113 - one global ordinal spans nested shapes
            fragments.extend(child_fragments)

        return ParsedDocument(
            parser_name=self.name,
            parser_version=self.version,
            fragments=fragments,
            assets=assets,
            metadata={"slide_count": len(presentation.slides)},
        )
